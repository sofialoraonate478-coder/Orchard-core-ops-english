#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script para generar un informe en Word del sistema Help Desk
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import os

def crear_informe():
    """Crea un informe IEEE 830 en formato Word"""
    
    # Crear documento
    doc = Document()
    
    # Título
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_run = titulo.add_run('INFORME IEEE 830')
    titulo_run.font.size = Pt(24)
    titulo_run.font.bold = True
    titulo_run.font.color.rgb = RGBColor(0, 0, 128)
    
    # Subtítulo
    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitulo_run = subtitulo.add_run('Sistema Help Desk - Especificación de Requisitos')
    subtitulo_run.font.size = Pt(14)
    subtitulo_run.font.color.rgb = RGBColor(64, 64, 64)
    
    # Fecha
    fecha = doc.add_paragraph()
    fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fecha_run = fecha.add_run(f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}')
    fecha_run.font.size = Pt(10)
    fecha_run.font.italic = True
    
    doc.add_paragraph()  # Espaciador
    
    # Sección 1: Introducción
    doc.add_heading('1. Introducción', level=1)
    doc.add_paragraph(
        'Este documento especifica los requisitos del sistema Help Desk, '
        'un aplicativo web para la gestión de tickets de soporte técnico.'
    )
    
    # Sección 2: Propósito
    doc.add_heading('2. Propósito', level=1)
    doc.add_paragraph(
        'El propósito de este documento es definir los requisitos funcionales y no funcionales '
        'para el sistema de gestión de Help Desk, siguiendo el estándar IEEE 830.'
    )
    
    # Sección 3: Características del Producto
    doc.add_heading('3. Características del Producto', level=1)
    
    features = [
        'Gestión de tickets de soporte',
        'Autenticación de usuarios',
        'Dashboard de reportes',
        'Base de datos MySQL',
        'Interfaz web responsiva',
    ]
    
    for feature in features:
        doc.add_paragraph(feature, style='List Bullet')
    
    # Sección 4: Requisitos Funcionales
    doc.add_heading('4. Requisitos Funcionales', level=1)
    
    doc.add_heading('4.1 Gestión de Tickets', level=2)
    doc.add_paragraph('El sistema debe permitir crear, leer, actualizar y eliminar tickets.')
    
    doc.add_heading('4.2 Autenticación', level=2)
    doc.add_paragraph('El sistema debe validar credenciales de usuario (email y contraseña).')
    
    doc.add_heading('4.3 Reportes', level=2)
    doc.add_paragraph('El sistema debe generar reportes de tickets resueltos y pendientes.')
    
    # Sección 5: Requisitos No Funcionales
    doc.add_heading('5. Requisitos No Funcionales', level=1)
    
    doc.add_heading('5.1 Rendimiento', level=2)
    doc.add_paragraph('El sistema debe responder en menos de 2 segundos a cualquier solicitud.')
    
    doc.add_heading('5.2 Seguridad', level=2)
    doc.add_paragraph('Las contraseñas deben almacenarse de forma segura con hash.')
    
    # Sección 6: Conclusión
    doc.add_heading('6. Conclusión', level=1)
    doc.add_paragraph(
        'Este documento define los requisitos para el sistema Help Desk. '
        'El desarrollo debe seguir estas especificaciones para garantizar la calidad del producto.'
    )
    
    # Guardar documento
    filename = 'INFORME_IEEE830_HELPDESK.docx'
    doc.save(filename)
    print(f'✓ Informe generado exitosamente: {filename}')
    print(f'✓ Ubicación: {os.path.abspath(filename)}')
    
    return filename

if __name__ == '__main__':
    try:
        crear_informe()
        print('✓ Proceso completado correctamente.')
    except Exception as e:
        print(f'✗ Error al generar el informe: {e}')
        exit(1)
